from __future__ import annotations

import io
import json
import os
import shutil
from typing import List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Alignment

from app.models.models import MetonymyEntry

TARGET_SHEET = "METONYMY CODING"
REQUIRED     = ("article_id", "expression", "type", "occurrence")


class WorkbookService:
    """
    Owns a single open workbook for the session.
    Provides: open / import / undo / export / sheet-data queries.
    """

    def __init__(self):
        self._wb            = None
        self._original_path = ""
        self._working_path  = ""
        self._snapshots:    list[bytes] = []   # for undo
        self._import_count  = 0
        self.unsaved        = False

        # Lists read from the workbook for validation
        self.valid_article_ids: set[str]  = set()
        self.valid_types:       list[str] = []

    # ── Open ──────────────────────────────────────────────────────────────────
    def open(self, path: str) -> dict:
        self._original_path = path
        base, ext = os.path.splitext(path)
        self._working_path  = base + "_working" + ext
        shutil.copy2(path, self._working_path)

        self._wb           = load_workbook(self._working_path, keep_vba=True)
        self._snapshots    = []
        self._import_count = 0
        self.unsaved       = False
        self._refresh_lists()
        return self.info()

    def is_open(self) -> bool:
        return self._wb is not None

    def _refresh_lists(self) -> None:
        if "LISTS & SETTINGS" not in self._wb.sheetnames:
            return
        ws = self._wb["LISTS & SETTINGS"]
        self.valid_article_ids = set()
        self.valid_types       = []
        for r in range(3, 500):
            a = ws.cell(r, 1).value
            if a: self.valid_article_ids.add(str(a).strip())
            t = ws.cell(r, 3).value
            if t: self.valid_types.append(str(t).strip())

    # ── Info & sheet data ─────────────────────────────────────────────────────
    def info(self) -> dict:
        if not self._wb:
            return {}
        data_rows = 0
        if TARGET_SHEET in self._wb.sheetnames:
            ws = self._wb[TARGET_SHEET]
            data_rows = sum(
                1 for row in ws.iter_rows(min_row=3, max_row=ws.max_row)
                if row[1].value
            )
        return {
            "filename":   os.path.basename(self._original_path),
            "path":       self._original_path,
            "sheets":     self._wb.sheetnames,
            "data_rows":  data_rows,
            "imports":    self._import_count,
        }

    def sheet_data(
        self, sheet_name: str, max_rows: int = 300, max_cols: int = 15
    ) -> Tuple[list, list]:
        if not self._wb or sheet_name not in self._wb.sheetnames:
            return [], []
        ws   = self._wb[sheet_name]
        rows = []
        for row in ws.iter_rows(
            min_row=1,
            max_row=min(ws.max_row, max_rows),
            max_col=min(ws.max_column, max_cols),
            values_only=True,
        ):
            rows.append(list(row))
        headers = rows[0] if rows else []
        return headers, rows

    # ── JSON parsing ──────────────────────────────────────────────────────────
    @staticmethod
    def parse_json(text: str) -> Tuple[List[dict], Optional[str]]:
        try:
            data = json.loads(text)
        except json.JSONDecodeError as e:
            return [], str(e)

        if isinstance(data, list):
            entries = data
        elif isinstance(data, dict):
            entries = data.get("entries", [])
            if not entries and any(k in data for k in ("article_id", "expression")):
                entries = [data]
        else:
            return [], "JSON must be an array or object with an 'entries' key."

        entries = [
            e for e in entries
            if isinstance(e, dict)
            and "total_unique_expressions" not in e
            and "frequency_by_type" not in e
        ]
        return entries, None

    def validate_entries(self, entries: List[dict]) -> List[str]:
        warnings = []
        for i, e in enumerate(entries, 1):
            missing = [f for f in REQUIRED if not e.get(f) and e.get(f) != 0]
            if missing:
                warnings.append(f"Entry {i}: missing {', '.join(missing)}")
            aid = str(e.get("article_id", "")).strip()
            if aid and self.valid_article_ids and aid not in self.valid_article_ids:
                warnings.append(
                    f"Entry {i}: article_id '{aid}' not found in ARTICLES list"
                )
        return warnings

    # ── Snapshot / undo ───────────────────────────────────────────────────────
    def _snapshot(self) -> None:
        buf = io.BytesIO()
        self._wb.save(buf)
        self._snapshots.append(buf.getvalue())
        if len(self._snapshots) > 20:
            self._snapshots.pop(0)

    def undo(self) -> Optional[str]:
        if not self._snapshots:
            return None
        data = self._snapshots.pop()
        self._wb = load_workbook(io.BytesIO(data), keep_vba=True)
        self._import_count = max(0, self._import_count - 1)
        self.unsaved = True
        self._wb.save(self._working_path)
        return f"Undone import #{self._import_count + 1}."

    # ── Import ────────────────────────────────────────────────────────────────
    def import_entries(self, entries: List[dict]) -> dict:
        if TARGET_SHEET not in self._wb.sheetnames:
            return {"error": f"Sheet '{TARGET_SHEET}' not found."}

        self._snapshot()
        ws       = self._wb[TARGET_SHEET]
        dest_row = self._next_empty_row(ws)
        written = skipped = 0
        warnings = []
        aln = Alignment(horizontal="right", vertical="center", wrap_text=True)

        for i, entry in enumerate(entries, 1):
            aid    = str(entry.get("article_id",         "")).strip()
            expr   = str(entry.get("expression",          "")).strip()
            etype  = str(entry.get("type",                "")).strip()
            occ    = entry.get("occurrence")
            ctx    = str(entry.get("context",              "")).strip()
            src    = str(entry.get("source_page",          "")).strip()
            lit    = str(entry.get("literal_referent",     "")).strip()
            met    = str(entry.get("metonymic_referent",   "")).strip()
            interp = str(entry.get("interpretation",       "")).strip()
            notes  = str(entry.get("notes",                "")).strip()

            missing = [
                f for f in REQUIRED
                if not entry.get(f) and entry.get(f) != 0
            ]
            if missing:
                warnings.append(
                    f"Entry {i}: missing {', '.join(missing)} — skipped."
                )
                skipped += 1
                continue

            r = dest_row
            # Auto-formulas (only write if cell is empty)
            if not ws.cell(r, 1).value:
                ws.cell(r, 1).value = (
                    f'=IF(B{r}="","","M"&TEXT(ROW()-2,"0000"))')
            if not ws.cell(r, 3).value:
                ws.cell(r, 3).value = (
                    f'=IFERROR(VLOOKUP(B{r},ARTICLES!$A:$D,4,FALSE),"")')
            if not ws.cell(r, 4).value:
                ws.cell(r, 4).value = (
                    f'=IFERROR(VLOOKUP(B{r},ARTICLES!$A:$B,2,FALSE),"")')

            ws.cell(r,  2).value = aid
            ws.cell(r,  5).value = expr
            ws.cell(r,  6).value = int(occ)
            ws.cell(r,  7).value = etype
            if ctx:    ws.cell(r,  8).value = ctx
            if src:    ws.cell(r,  9).value = src
            if lit:    ws.cell(r, 10).value = lit
            if met:    ws.cell(r, 11).value = met
            if interp: ws.cell(r, 12).value = interp
            if notes:  ws.cell(r, 13).value = notes

            for col in range(1, 14):
                c = ws.cell(r, col)
                if c.value:
                    c.alignment = aln

            dest_row += 1
            written  += 1

        self._wb.save(self._working_path)
        self._import_count += 1
        self.unsaved = True

        total_rows = sum(
            1 for row in ws.iter_rows(min_row=3, max_row=ws.max_row)
            if row[1].value
        )
        return {
            "import_number":   self._import_count,
            "written":         written,
            "skipped":         skipped,
            "warnings":        warnings,
            "total_data_rows": total_rows,
            "sheet":           TARGET_SHEET,
        }

    @staticmethod
    def _next_empty_row(ws) -> int:
        dest = 3
        for row in ws.iter_rows(min_row=3, max_row=5000, min_col=2, max_col=2):
            if row[0].value and str(row[0].value).strip():
                dest = row[0].row + 1
            else:
                break
        return dest

    # ── Export ────────────────────────────────────────────────────────────────
    def export(self, dest_path: str) -> None:
        self._wb.save(dest_path)
        self.unsaved = False

    def save_working(self) -> None:
        if self._wb and self._working_path:
            self._wb.save(self._working_path)
            self.unsaved = False

    @property
    def original_path(self) -> str:
        return self._original_path

    @property
    def import_count(self) -> int:
        return self._import_count
